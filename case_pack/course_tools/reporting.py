"""Supplied helpers for D1-S07 (filling the management template) and D1-S08 (the monthly run).

Learners call these from their scripts. They do not need to read this file.
The helpers reuse the checked Northbridge tools in northbridge.py; messages are plain English.
"""
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from case_pack.course_tools import northbridge as nb

TEMPLATE = Path("case_pack/templates/northbridge_reporting_template.xlsx")
BYLINE_COLUMNS = ["period", "entity", "entity_name", "report_line", "report_line_label",
                  "actual_amount", "budget_amount", "favourable_variance"]
BLOCKED_TEXT = "No figures: the release is blocked. See the Controls and Exceptions sheets."
_TEMPLATE_FINGERPRINT = None


# ---------------------------------------------------------------------------------------------- D1-S08 steps
def load_month(input_folder, period, mapping_file=None):
    """Step: read the files the month's manifest lists, plus mapping, budget and the units' own totals."""
    inputs = nb.load_inputs(input_folder, period, mapping_file=mapping_file)
    received = ", ".join(inputs["received_files"]) or "none"
    print(f"Loaded {len(inputs['transactions'])} transaction rows for {period} from: {received}")
    if inputs["missing_files"]:
        print("Listed in the manifest but NOT received:", ", ".join(inputs["missing_files"]))
    print("Account mapping used:", inputs["mapping_name"])
    return inputs


def run_checks(inputs):
    """Step: run the nine controls. Returns (control results, exceptions, release decision)."""
    controls, exceptions, release = nb.run_controls(inputs)
    for _, check in controls.iterrows():
        print(f"  {check['check_id']} {check['check_name']:<29} {check['status']}")
    print(f"Release decision: {release} ({len(exceptions)} exception rows)")
    return controls, exceptions, release


def save_check_results(controls, exceptions, output_folder):
    """Step: always save the control results and exceptions, whatever the decision."""
    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)
    nb.write_csv(controls, output_folder / "control_summary.csv")
    nb.write_csv(exceptions, output_folder / "exceptions.csv")
    print(f"Saved control_summary.csv and exceptions.csv in {output_folder.as_posix()}")


def build_figures(inputs):
    """Step: calculate the nine report lines and the unit summary. Only call this when the release is not blocked."""
    entities = [entity for entity, _ in nb.expected_exports(inputs)]
    try:
        actuals = nb.build_actuals(inputs["transactions"], inputs["account_mapping"], inputs["period"], entities)
    except ValueError as problem:
        raise SystemExit(f"STOPPED: the figures cannot be calculated safely ({problem}).\n"
                         "This happens when figures are built for a month whose checks failed. "
                         "Only call build_figures when the release is not blocked.")
    summary = nb.build_summary(actuals, inputs["budget"], inputs["entities"], inputs["period"])
    print(f"Calculated {len(actuals)} report lines for {inputs['period']}")
    return summary


def save_figures(summary, output_folder):
    """Step: save the nine report lines as actuals_by_line.csv."""
    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)
    table = summary["by_line"][["period", "entity", "report_line", "actual_amount"]].copy()
    table["actual_amount"] = table["actual_amount"].map(nb.money_text)
    nb.write_csv(table, output_folder / "actuals_by_line.csv")
    print(f"Saved actuals_by_line.csv in {output_folder.as_posix()}")


def write_pack(inputs, controls, exceptions, release, summary, output_folder):
    """Step: fill a copy of the management template and save it. Figures are left out when blocked."""
    path = nb.write_workbook(output_folder, inputs, release, controls, exceptions, summary, overwrite=True)
    print(f"Saved workbook {path.as_posix()}")
    return path


# ---------------------------------------------------------------------------------------------- D1-S07 helpers
def prepare_pack(case):
    """Run the checks for December and return everything the template needs.

    case is "clean" (all checks pass) or "blocked" (a mapping lists account 0500 twice).
    """
    folders = {"clean": "case_pack/data/clean", "blocked": "case_pack/data/faulty/F03_duplicate_mapping"}
    if case not in folders:
        raise SystemExit('CASE must be "clean" or "blocked" (with the quotation marks).')
    inputs = nb.load_inputs(folders[case], "2025-12")
    controls, exceptions, release = nb.run_controls(inputs)
    summary = build_figures(inputs) if release != "blocked" else None
    by_line = None
    if summary is not None:
        by_line = summary["by_line"][BYLINE_COLUMNS].copy()
        money_columns = ["actual_amount", "budget_amount", "favourable_variance"]
        by_line[money_columns] = by_line[money_columns].round(2)
    failed = controls.loc[controls["status"] != "pass", "check_id"].tolist()
    print(f"Case: {case}. Checks not passing: {', '.join(failed) or 'none'}. Release decision: {release}")
    return {"case": case, "inputs": inputs, "controls": controls, "exceptions": exceptions,
            "release": release, "summary": summary, "by_line": by_line}


def open_template():
    """Open the master template in memory. The master file itself is never saved over."""
    global _TEMPLATE_FINGERPRINT
    _TEMPLATE_FINGERPRINT = TEMPLATE.stat().st_size, round(TEMPLATE.stat().st_mtime, 3)
    return load_workbook(TEMPLATE)


def fill_other_sheets(workbook, pack):
    """Fill Summary, Controls, Exceptions and ReadMe. (You fill ByLine.)"""
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    inputs = pack["inputs"]
    nb.fill_summary_sheet(workbook["Summary"], inputs["period"], pack["release"], generated,
                          nb.REFERENCE_PACK_VERSION, pack["summary"])
    if pack["summary"] is None:
        # A blocked pack has no figures, so the cross-check formula would read 0.00: the same value that means
        # "the figures agree" in a released pack.
        workbook["Summary"]["C14"] = "Not applicable: figures withheld"
    nb.fill_controls_sheet(workbook["Controls"], pack["controls"])
    nb.fill_exceptions_sheet(workbook["Exceptions"], pack["exceptions"])
    name = nb.workbook_filename(inputs["period"], 1, pack["release"])
    nb.fill_readme_sheet(workbook["ReadMe"], nb.readme_lines(inputs, pack["release"], pack["controls"], generated,
                                                             nb.REFERENCE_PACK_VERSION, name))


def save_new_workbook(workbook, pack, output_folder):
    """Save under a new dated name in the output folder, replacing only this exercise's earlier output."""
    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)
    path = output_folder / nb.workbook_filename(pack["inputs"]["period"], 1, pack["release"])
    if path.resolve() == TEMPLATE.resolve():
        raise SystemExit("Refusing to save over the master template.")
    try:
        workbook.save(path)
    except PermissionError:
        raise SystemExit(f"Could not save {path.name}: it is probably open in Excel. Close it and run again.")
    print(f"Saved {path.as_posix()}")
    return path


def check_byline(path, pack):
    """Reopen the saved file and check the ByLine sheet you filled."""
    if _TEMPLATE_FINGERPRINT and _TEMPLATE_FINGERPRINT != (TEMPLATE.stat().st_size, round(TEMPLATE.stat().st_mtime, 3)):
        print("PROBLEM: the master template itself was changed. Your code must write into the supplied `sheet` and "
              "let the supplied save step write into outputs/. Restore "
              "case_pack/templates/northbridge_reporting_template.xlsx from the course download.")
    book = load_workbook(path)
    sheet = book["ByLine"]
    written = [[sheet.cell(row, column).value for column in range(1, 9)] for row in range(2, 12)]
    figures = [row for row in written if isinstance(row[5], (int, float))]
    # Look for the same nine amounts written somewhere other than rows 2 to 10.
    elsewhere = [r for r in range(12, sheet.max_row + 1)
                 if isinstance(sheet.cell(r, 6).value, (int, float)) or isinstance(sheet.cell(r, 2).value, str)
                 and str(sheet.cell(r, 2).value).startswith("N0")]
    if not figures and elsewhere:
        print(f"PROBLEM: nine rows were written, but starting at row {elsewhere[0]} instead of row 2. "
              "The loop must write to row 2 + the row number, not append to the end of the sheet.")
        return
    if any(isinstance(row[0], (int, float)) for row in written):
        print("PROBLEM: column A holds a number, not the period (2025-12): the loop is writing the table's row "
              "numbers (its index) as the first column, so every amount is one column too far right. "
              "Ask Copilot for itertuples(index=False), or to leave the index out.")
        return
    definitions_kept = str(sheet.cell(13, 1).value or "").startswith("Definitions")
    if pack["release"] == "blocked":
        ok = not figures and sheet["A2"].value == BLOCKED_TEXT
        print("OK: the blocked pack shows no figures on ByLine." if ok else
              "PROBLEM: a blocked pack must not show figures. ByLine should only say the release is blocked.")
    else:
        if not figures:
            print("NOT DONE YET: ByLine has no figures in rows 2 to 10. If you did paste a loop, check that it writes "
                  "into the supplied `sheet` (it must not create a new workbook) and starts at row 2.")
            return
        table = pd.DataFrame(figures, columns=BYLINE_COLUMNS)
        checkpoint = pd.read_csv("case_pack/data/clean/checkpoints/actuals_by_line_2025-12.csv")
        merged = checkpoint.merge(table, on=["entity", "report_line"], how="left", suffixes=("_checkpoint", ""))
        matches = (merged["actual_amount"] - merged["actual_amount_checkpoint"]).abs().le(0.01).sum()
        print(f"OK: 9 lines written starting at row 2." if len(figures) == 9 else
              f"PROBLEM: {len(figures)} lines with figures; expected 9 starting at row 2.")
        print(f"OK: all 9 actual amounts match the checkpoint." if matches == 9 else
              f"PROBLEM: {matches} of 9 actual amounts match the checkpoint. Check the column order A to H.")
        in_order = all(row[0] == "2025-12" and row[1] in ("N01", "N02", "N03") for row in figures)
        print("OK: column A is the period and column B the unit." if in_order else
              "PROBLEM: columns are in the wrong order. A = period, B = entity, ... H = favourable_variance.")
    print("OK: the definitions block below the table is untouched." if definitions_kept else
          "PROBLEM: the definitions text that starts at row 13 was overwritten. Write only rows 2 to 10.")


# ---------------------------------------------------------------------------------------------- D1-S08 run support
EXERCISE_FILES = ("control_summary.csv", "exceptions.csv", "actuals_by_line.csv")


def fresh_output_folder(output_folder):
    """Remove this exercise's own earlier output files, so old figures can never sit beside a new result."""
    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)
    for path in list(output_folder.glob("*.xlsx")) + [output_folder / name for name in EXERCISE_FILES]:
        if path.exists():
            try:
                path.unlink()
            except PermissionError:
                raise SystemExit(f"Could not replace {path.name}: it is probably open in Excel. Close it and run again.")
    return output_folder


def check_month_run(output_folder, period, release):
    """Check the files a monthly run must leave behind, for a blocked and for a released month."""
    output_folder = Path(output_folder)
    if release is None:
        print("NOT DONE YET: run_month returned nothing. Finish YOUR TURN and make it return the release decision.")
        return
    has = {name: (output_folder / name).exists() for name in EXERCISE_FILES}
    books = sorted(p.name for p in output_folder.glob("*.xlsx"))
    print("OK: control_summary.csv and exceptions.csv were saved." if has["control_summary.csv"] and has["exceptions.csv"]
          else "PROBLEM: control_summary.csv and exceptions.csv must always be saved, even when the month is blocked.")
    if release == "blocked":
        print("OK: no figures file was saved for a blocked month." if not has["actuals_by_line.csv"]
              else "PROBLEM: actuals_by_line.csv was saved although the release is blocked. Only build figures when not blocked.")
        print("OK: the workbook is marked _BLOCKED." if books and all("_BLOCKED" in b for b in books)
              else "PROBLEM: expected one workbook whose name ends _BLOCKED. Did you call write_pack?")
        return
    if not has["actuals_by_line.csv"]:
        print("PROBLEM: the release is not blocked, but actuals_by_line.csv is missing. Call build_figures and save_figures.")
        return
    mine = pd.read_csv(output_folder / "actuals_by_line.csv", dtype={"period": str})
    checkpoint = pd.read_csv(f"case_pack/data/clean/checkpoints/actuals_by_line_{period}.csv", dtype={"period": str})
    merged = checkpoint.merge(mine, on=["period", "entity", "report_line"], how="left", suffixes=("_checkpoint", ""))
    matches = int((merged["actual_amount"] - merged["actual_amount_checkpoint"]).abs().le(0.01).sum())
    print(f"OK: all 9 figures match the {period} checkpoint." if matches == 9 and len(mine) == 9
          else f"PROBLEM: {matches} of 9 figures match the {period} checkpoint.")
    print("OK: a released workbook was saved." if books and not any("_BLOCKED" in b for b in books)
          else "PROBLEM: expected one workbook without _BLOCKED in its name. Did you pass the figures to write_pack?")
