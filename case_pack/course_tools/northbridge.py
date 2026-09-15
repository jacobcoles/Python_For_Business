"""Supplied Northbridge reporting tools. No learner edits are needed here."""
import argparse
import json
import re
import sys
from copy import copy
from datetime import datetime, timezone
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Alignment
from openpyxl.worksheet.properties import PageSetupProperties
TEMPLATE_PATH = Path('case_pack/templates/northbridge_reporting_template.xlsx')
DEFAULT_OUTPUT_FOLDER = Path('outputs')
REPORT_LINES = ['revenue', 'direct_cost', 'overhead']
REPORT_LINE_LABELS = {'revenue': 'Revenue', 'direct_cost': 'Direct cost', 'overhead': 'Overhead'}
TRANSACTION_COLUMNS = ['transaction_id', 'period', 'entity', 'account_code', 'amount', 'currency', 'status']
BUSINESS_KEY = ['entity', 'period', 'transaction_id']
ALLOWED_STATUSES = {'posted', 'cancelled'}
RECONCILIATION_TOLERANCE_CENTS = 1
MOVEMENT_AMOUNT_THRESHOLD_CENTS = 10000
MOVEMENT_PCT_THRESHOLD = Decimal('0.20')
CHECKS = {'C01': ('completeness_files', 'critical', True), 'C02': ('completeness_rows', 'critical', True), 'C03': ('duplicate_business_key', 'critical', True), 'C04': ('mapping_cardinality', 'critical', True), 'C05': ('unmapped_accounts', 'critical', True), 'C06': ('required_values', 'critical', True), 'C07': ('currency_supported', 'critical', True), 'C08': ('reconciliation_posted_amount', 'critical', True), 'C09': ('movement_review', 'review', True)}
SUMMARY_COLUMNS = ['check_id', 'check_name', 'severity', 'scope', 'expected', 'observed', 'tolerance', 'status', 'affected_count', 'detail_ref']
EXCEPTION_COLUMNS = ['check_id', 'severity', 'source_file', 'source_row', 'entity', 'period', 'transaction_id', 'issue', 'observed_value', 'expected_value', 'action_note']
TRANSACTION_FILENAME = re.compile('transactions_(\\d{4}-\\d{2})_(N\\d{2})\\.csv')
PLAIN_NUMBER = '-?\\d+(\\.\\d+)?'

def to_cents(value):
    return int((Decimal(repr(float(value))) * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))

def money(value):
    return float(Decimal(repr(float(value))).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)) + 0.0

def money_text(value, signed=False):
    text = f'{Decimal(repr(float(value))).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)}'
    if signed and (not text.startswith('-')):
        text = '+' + text
    return '0.00' if text in ('-0.00', '+-0.00') else text

def parse_amount(text_values):
    text = text_values.astype(str).str.strip()
    looks_like_a_number = text.str.fullmatch(PLAIN_NUMBER)
    return pd.to_numeric(text.where(looks_like_a_number), errors='coerce')

def is_blank(text_values):
    return text_values.astype(str).str.strip() == ''

def previous_period(period):
    year, month = (int(period[:4]), int(period[5:7]))
    return f'{year - 1}-12' if month == 1 else f'{year}-{month - 1:02d}'

def read_transaction_file(path):
    rows = pd.read_csv(path, dtype=str, keep_default_na=False)
    missing_columns = [c for c in TRANSACTION_COLUMNS if c not in rows.columns]
    if missing_columns:
        raise ValueError(f'{path.name}: export is missing columns {missing_columns}')
    rows = rows[TRANSACTION_COLUMNS].copy()
    rows['source_file'] = path.name
    rows['source_row'] = range(1, len(rows) + 1)
    return rows

def load_transaction_files(input_folder, relative_names, period):
    frames, missing = ([], [])
    for relative_name in relative_names:
        path = input_folder / relative_name
        match = TRANSACTION_FILENAME.fullmatch(path.name)
        if not match or match.group(1) != period:
            raise ValueError(f'Input manifest lists {relative_name!r}, which is not a transactions_{period}_<unit>.csv export for period {period}')
        if path.exists():
            frames.append(read_transaction_file(path))
        else:
            missing.append(path.name)
    if frames:
        rows = pd.concat(frames, ignore_index=True)
    else:
        rows = pd.DataFrame(columns=TRANSACTION_COLUMNS + ['source_file', 'source_row'])
    return (rows, missing)

def read_reference_table(input_folder, manifest, file_name, override_path=None):
    if override_path is not None:
        path = Path(override_path)
    else:
        listed = [name for name in manifest.get('reference_files', []) if Path(name).name == file_name]
        if not listed:
            return (None, file_name)
        path = input_folder / listed[0]
    if not path.exists():
        return (None, path.name)
    return (pd.read_csv(path, dtype=str, keep_default_na=False), path.name)

def load_inputs(input_folder, period, manifest_path=None, mapping_file=None):
    input_folder = Path(input_folder)
    manifest_path = Path(manifest_path) if manifest_path else input_folder / f'input_manifest_{period}.json'
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    if manifest.get('reporting_period') != period:
        raise ValueError(f'{manifest_path.name} is for {manifest.get('reporting_period')!r}, not the requested period {period!r}')
    transactions, missing_files = load_transaction_files(input_folder, manifest['transaction_files'], period)
    account_mapping, mapping_name = read_reference_table(input_folder, manifest, 'account_mapping.csv', mapping_file)
    entities, _ = read_reference_table(input_folder, manifest, 'entities.csv')
    budget, _ = read_reference_table(input_folder, manifest, 'budget.csv')
    source_controls, _ = read_reference_table(input_folder, manifest, 'source_controls.csv')
    expected_files, _ = read_reference_table(input_folder, manifest, 'expected_files.csv')
    prior_period = manifest.get('prior_period') or previous_period(period)
    prior_actuals, prior_transactions, prior_missing_files, prior_source = (None, None, [], None)
    if manifest.get('prior_period_actuals'):
        path = input_folder / manifest['prior_period_actuals']
        prior_source = path.name
        if path.exists():
            prior_actuals = pd.read_csv(path, dtype=str, keep_default_na=False)
            prior_actuals['actual_amount'] = parse_amount(prior_actuals['actual_amount'])
        else:
            prior_missing_files = [path.name]
    elif manifest.get('prior_period_transaction_files'):
        prior_transactions, prior_missing_files = load_transaction_files(input_folder, manifest['prior_period_transaction_files'], prior_period)
        prior_source = ', '.join((Path(n).name for n in manifest['prior_period_transaction_files']))
    listed_files = [Path(n).name for n in manifest['transaction_files']]
    return {'input_folder': input_folder, 'manifest_name': manifest_path.name, 'pack_version': manifest.get('pack_version'), 'period': period, 'prior_period': prior_period, 'transactions': transactions, 'listed_files': listed_files, 'received_files': [name for name in listed_files if name not in missing_files], 'missing_files': missing_files, 'account_mapping': account_mapping, 'mapping_name': mapping_name, 'entities': entities, 'budget': budget, 'source_controls': source_controls, 'expected_files': expected_files, 'prior_actuals': prior_actuals, 'prior_transactions': prior_transactions, 'prior_missing_files': prior_missing_files, 'prior_source': prior_source}

def summary_row(check_id, scope, expected, observed, tolerance, status, affected_count, has_detail):
    name, severity, _ = CHECKS[check_id]
    return {'check_id': check_id, 'check_name': name, 'severity': severity, 'scope': scope, 'expected': expected, 'observed': observed, 'tolerance': tolerance, 'status': status, 'affected_count': int(affected_count), 'detail_ref': f'exceptions.csv check_id={check_id}' if has_detail else ''}

def exception_row(check_id, issue, observed_value, expected_value, action_note, source_file='', source_row='', entity='', period='', transaction_id=''):
    return {'check_id': check_id, 'severity': CHECKS[check_id][1], 'source_file': source_file, 'source_row': source_row, 'entity': entity, 'period': period, 'transaction_id': transaction_id, 'issue': issue, 'observed_value': observed_value, 'expected_value': expected_value, 'action_note': action_note}

def exception_from_source_row(check_id, row, issue, observed_value, expected_value, action_note):
    return exception_row(check_id, issue, observed_value, expected_value, action_note, source_file=row['source_file'], source_row=int(row['source_row']), entity=row['entity'], period=row['period'], transaction_id=row['transaction_id'])

def combine_scope_statuses(statuses):
    if 'fail' in statuses:
        return 'fail'
    if 'not_run' in statuses:
        return 'not_run'
    if 'warning' in statuses:
        return 'warning'
    return 'pass'

def expected_exports(inputs):
    expected = inputs['expected_files']
    if expected is not None:
        this_period = expected[expected['period'] == inputs['period']]
        return list(zip(this_period['entity'], this_period['filename']))
    return [(TRANSACTION_FILENAME.fullmatch(name).group(2), name) for name in inputs['listed_files']]

def check_c01_completeness_files(inputs):
    period = inputs['period']
    if inputs['expected_files'] is None:
        return (summary_row('C01', period, 'expected_files.csv available', 'expected_files.csv missing or unreadable', 'n/a', 'not_run', 0, False), [])
    expected = expected_exports(inputs)
    if not expected:
        return (summary_row('C01', period, f'expected files listed for {period}', f'expected_files.csv lists no files for {period}', 'n/a', 'not_run', 0, False), [])
    missing = [(entity, name) for entity, name in expected if name not in inputs['received_files']]
    exceptions = [exception_row('C01', 'expected export not received', 'file not present in the input', name, f'Obtain {name} from unit {entity}; do not report {period} without it.', source_file=name, entity=entity, period=period) for entity, name in missing]
    observed = f'{len(expected) - len(missing)} of {len(expected)} present'
    if missing:
        observed += '; missing: ' + ', '.join((name for _, name in missing))
    return (summary_row('C01', period, f'{len(expected)} expected files present', observed, 'n/a', 'fail' if missing else 'pass', len(missing), bool(exceptions)), exceptions)

def control_total_row(inputs, entity):
    controls = inputs['source_controls']
    if controls is None:
        return None
    match = controls[(controls['period'] == inputs['period']) & (controls['entity'] == entity)]
    return None if match.empty else match.iloc[0]

def check_c02_completeness_rows(inputs):
    period, rows = (inputs['period'], inputs['transactions'])
    statuses, expected_parts, observed_parts, exceptions, failing = ([], [], [], [], 0)
    for entity, filename in expected_exports(inputs):
        control = control_total_row(inputs, entity)
        if filename not in inputs['received_files']:
            statuses.append('not_run')
            observed_parts.append(f'{entity}: not_run - source file missing - see C01')
            continue
        if control is None:
            statuses.append('not_run')
            observed_parts.append(f'{entity}: not_run - no control row in source_controls.csv')
            continue
        unit_rows = rows[rows['source_file'] == filename]
        raw_count = len(unit_rows)
        posted_count = int((unit_rows['status'] == 'posted').sum())
        expected_raw, expected_posted = (int(control['expected_raw_rows']), int(control['expected_posted_rows']))
        expected_parts.append(f'{entity}: raw {expected_raw}, posted {expected_posted}')
        observed_parts.append(f'{entity}: raw {raw_count}, posted {posted_count}')
        if raw_count == expected_raw and posted_count == expected_posted:
            statuses.append('pass')
        else:
            statuses.append('fail')
            failing += 1
            exceptions.append(exception_row('C02', "row count differs from the unit's control total", f'raw {raw_count}; posted {posted_count}', f'raw {expected_raw}; posted {expected_posted}', "Find the extra or missing rows before using this unit's figures.", source_file=filename, entity=entity, period=period))
    status = combine_scope_statuses(statuses) if statuses else 'not_run'
    return (summary_row('C02', period, '; '.join(expected_parts) or 'control rows available', '; '.join(observed_parts) or 'no expected exports', 'exact', status, failing, bool(exceptions)), exceptions)

def check_c03_duplicate_business_key(inputs, duplicate_finder=None):
    period, rows = (inputs['period'], inputs['transactions'])
    keyed = rows[~(is_blank(rows['entity']) | is_blank(rows['period']) | is_blank(rows['transaction_id']))]
    duplicates = duplicate_finder(keyed) if duplicate_finder else keyed[keyed.duplicated(subset=BUSINESS_KEY, keep=False)]
    if duplicates is None:
        return summary_row("C03", period, "completed duplicate check", "not completed", "exact", "not_run", 0, False), []
    duplicated_keys = duplicates[BUSINESS_KEY].drop_duplicates()
    exceptions = []
    for _, row in duplicates.iterrows():
        key = f'{row['entity']}|{row['period']}|{row['transaction_id']}'
        exceptions.append(exception_from_source_row('C03', row, 'duplicate business key (entity, period, transaction_id)', key, 'each key appears once', 'Ask the unit which occurrence is genuine; never drop duplicates to make counts match.'))
    observed = f'{len(duplicated_keys)} duplicated key(s), {len(duplicates)} row(s)'
    return (summary_row('C03', period, '0 duplicated keys', observed, 'exact', 'fail' if len(duplicated_keys) else 'pass', len(duplicated_keys), bool(exceptions)), exceptions)

def check_c04_mapping_cardinality(inputs):
    period, mapping = (inputs['period'], inputs['account_mapping'])
    if mapping is None:
        return (summary_row('C04', period, 'account_mapping.csv available', 'account mapping missing', 'n/a', 'not_run', 0, False), [])
    counts = mapping['account_code'].value_counts()
    offending_codes = sorted(counts[counts > 1].index)
    exceptions = []
    for position, row in mapping.iterrows():
        if row['account_code'] in offending_codes:
            exceptions.append(exception_row('C04', 'account code mapped more than once', f'{row['account_code']} -> {row['report_line']} ({row['account_name']})', 'exactly one mapping row per account_code', 'Ask finance which mapping is correct; do not join until the mapping is one-to-one.', source_file=inputs['mapping_name'], source_row=position + 1))
    observed = f'{len(offending_codes)} code(s) with more than one row'
    if offending_codes:
        observed += ': ' + ', '.join(offending_codes)
    return (summary_row('C04', period, '1 row per account_code', observed, 'exact', 'fail' if offending_codes else 'pass', len(offending_codes), bool(exceptions)), exceptions)

def check_c05_unmapped_accounts(inputs):
    period, rows, mapping = (inputs['period'], inputs['transactions'], inputs['account_mapping'])
    if mapping is None:
        return (summary_row('C05', period, 'account_mapping.csv available', 'account mapping missing', 'n/a', 'not_run', 0, False), [])
    known_codes = set(mapping['account_code'])
    posted = rows[rows['status'] == 'posted']
    unmapped = posted[~is_blank(posted['account_code']) & ~posted['account_code'].isin(known_codes)]
    exceptions = [exception_from_source_row('C05', row, 'posted row uses an account code that is not in the mapping', f'account_code={row['account_code']}; amount={row['amount']}', f'an account_code listed in {inputs['mapping_name']}', 'Ask finance for the mapping of this code; do not edit the export, bucket or drop the amount.') for _, row in unmapped.iterrows()]
    unmapped_amount = parse_amount(unmapped['amount']).sum()
    observed = f'{len(unmapped)} unmapped posted row(s), EUR {money_text(unmapped_amount)}'
    return (summary_row('C05', period, '0 unmapped posted rows', observed, 'exact', 'fail' if len(unmapped) else 'pass', len(unmapped), bool(exceptions)), exceptions)

def required_value_problems(row):
    problems = []
    for column in TRANSACTION_COLUMNS:
        if str(row[column]).strip() == '':
            problems.append(f'{column} is blank')
    file_period, file_entity = TRANSACTION_FILENAME.fullmatch(row['source_file']).groups()
    if str(row['amount']).strip() and (not re.fullmatch(PLAIN_NUMBER, str(row['amount']).strip())):
        problems.append(f'amount {row['amount']!r} is not a number')
    if str(row['status']).strip() and row['status'] not in ALLOWED_STATUSES:
        problems.append(f'status {row['status']!r} is not posted or cancelled')
    if str(row['period']).strip() and row['period'] != file_period:
        problems.append(f'period {row['period']!r} does not match file period {file_period}')
    if str(row['entity']).strip() and row['entity'] != file_entity:
        problems.append(f'entity {row['entity']!r} does not match file entity {file_entity}')
    return problems

def check_c06_required_values(inputs):
    period, rows = (inputs['period'], inputs['transactions'])
    exceptions = []
    for _, row in rows.iterrows():
        problems = required_value_problems(row)
        if problems:
            exceptions.append(exception_from_source_row('C06', row, 'required value missing or invalid: ' + '; '.join(problems), '; '.join((f'{c}={row[c]!r}' for c in TRANSACTION_COLUMNS)), 'all seven values present; amount a number; status posted or cancelled; period and entity match the file name', 'Return the row to the unit for correction; never default a blank to 0 or to posted.'))
    return (summary_row('C06', period, '0 rows with missing or invalid required values', f'{len(exceptions)} row(s) with missing or invalid required values', 'exact', 'fail' if exceptions else 'pass', len(exceptions), bool(exceptions)), exceptions)

def check_c07_currency_supported(inputs):
    period, rows = (inputs['period'], inputs['transactions'])
    wrong = rows[~is_blank(rows['currency']) & (rows['currency'] != 'EUR')]
    exceptions = [exception_from_source_row('C07', row, 'unsupported currency', f'currency={row['currency']}; amount={row['amount']}', 'EUR', f'Return the row to unit {row['entity']}; no conversion is attempted.') for _, row in wrong.iterrows()]
    return (summary_row('C07', period, 'all rows EUR', f'{len(wrong)} non-EUR row(s)', 'exact', 'fail' if len(wrong) else 'pass', len(wrong), bool(exceptions)), exceptions)

def check_c08_reconciliation(inputs, comparator=None):
    period, rows = (inputs['period'], inputs['transactions'])
    statuses, expected_parts, observed_parts, exceptions, failing = ([], [], [], [], 0)
    for entity, filename in expected_exports(inputs):
        control = control_total_row(inputs, entity)
        if filename not in inputs['received_files']:
            statuses.append('not_run')
            observed_parts.append(f'{entity}: not_run - source file missing - see C01')
            continue
        if control is None:
            statuses.append('not_run')
            observed_parts.append(f'{entity}: not_run - no control row in source_controls.csv')
            continue
        unit_rows = rows[rows['source_file'] == filename]
        amounts = parse_amount(unit_rows['amount'])
        unusable_amounts = int(amounts.isna().sum())
        unusable_statuses = int((~unit_rows['status'].isin(ALLOWED_STATUSES)).sum())
        if unusable_amounts or unusable_statuses:
            statuses.append('not_run')
            observed_parts.append(f'{entity}: not_run - {unusable_amounts} unusable amount(s) and {unusable_statuses} unusable status(es), total cannot be formed - see C06')
            continue
        received_total = amounts[unit_rows['status'] == 'posted'].sum()
        expected_total = float(control['expected_posted_amount'])
        difference_cents = to_cents(received_total) - to_cents(expected_total)
        expected_parts.append(f'{entity}: {money_text(expected_total)}')
        observed_parts.append(f'{entity}: {money_text(received_total)} (difference {money_text(difference_cents / 100, signed=True)})')
        learner_answer = comparator(received_total, expected_total) if comparator else None
        if comparator and learner_answer not in ("pass", "fail"):
            if learner_answer is not None:
                print(f"  NOTE: your reconciliation rule returned {learner_answer!r}, not 'pass' or 'fail', so C08 could not run. See D1-S06 Step 3.")
            statuses.append("not_run")
            observed_parts.append("not completed")
            continue
        if (learner_answer == "pass") if comparator else (abs(difference_cents) <= RECONCILIATION_TOLERANCE_CENTS):
            statuses.append('pass')
        else:
            statuses.append('fail')
            failing += 1
            exceptions.append(exception_row('C08', f"posted amount differs from the unit's control total by {money_text(difference_cents / 100, signed=True)}", money_text(received_total), money_text(expected_total), 'Find the rows behind the difference with the unit before using its figures.', source_file=filename, entity=entity, period=period))
    status = combine_scope_statuses(statuses) if statuses else 'not_run'
    return (summary_row('C08', period, '; '.join(expected_parts) or 'control rows available', '; '.join(observed_parts) or 'no expected exports', 'EUR 0.01', status, failing, bool(exceptions)), exceptions)

def prior_actuals_for_review(inputs, entities):
    if inputs['prior_missing_files']:
        return (None, f'prior-period input missing: {', '.join(inputs['prior_missing_files'])}')
    if inputs['prior_actuals'] is not None:
        prior = inputs['prior_actuals']
        prior = prior[prior['period'] == inputs['prior_period']]
        if prior['actual_amount'].isna().any():
            return (None, 'prior-period actuals contain a blank or unparseable amount')
        return (prior, None)
    prior_rows = inputs['prior_transactions']
    if prior_rows is None:
        return (None, 'prior-period actuals unavailable (the manifest supplies none)')
    known_codes = set(inputs['account_mapping']['account_code'])
    posted = prior_rows[prior_rows['status'] == 'posted']
    unusable = (~prior_rows['status'].isin(ALLOWED_STATUSES)).sum() + parse_amount(posted['amount']).isna().sum() + (posted['currency'] != 'EUR').sum() + (~posted['account_code'].isin(known_codes)).sum()
    if unusable:
        return (None, f'prior-period exports contain {int(unusable)} unusable value(s)')
    return (build_actuals(prior_rows, inputs['account_mapping'], inputs['prior_period'], entities), None)

def review_movements(current_actuals, prior_actuals, entities):
    results = []
    for entity in entities:
        for line in REPORT_LINES:
            current = current_actuals[(current_actuals['entity'] == entity) & (current_actuals['report_line'] == line)]
            prior = prior_actuals[(prior_actuals['entity'] == entity) & (prior_actuals['report_line'] == line)]
            result = {'entity': entity, 'report_line': line, 'current': None, 'prior': None, 'movement': None, 'movement_pct': None, 'outcome': 'not_run'}
            if len(current) != 1 or len(prior) != 1:
                results.append(result)
                continue
            current_cents = to_cents(current['actual_amount'].iloc[0])
            prior_cents = to_cents(prior['actual_amount'].iloc[0])
            movement_cents = current_cents - prior_cents
            result.update(current=current_cents / 100, prior=prior_cents / 100, movement=movement_cents / 100)
            if prior_cents == 0:
                result['movement_pct'] = 'not_comparable'
                big_new_activity = abs(current_cents) > MOVEMENT_AMOUNT_THRESHOLD_CENTS
                result['outcome'] = 'new_activity' if big_new_activity else 'silent'
            else:
                pct = Decimal(movement_cents) / Decimal(abs(prior_cents))
                result['movement_pct'] = pct
                big_amount = abs(movement_cents) > MOVEMENT_AMOUNT_THRESHOLD_CENTS
                big_percent = abs(pct) > MOVEMENT_PCT_THRESHOLD
                result['outcome'] = 'warning' if big_amount and big_percent else 'silent'
            results.append(result)
    return results

def pct_text(pct):
    if pct == 'not_comparable' or pct is None:
        return 'not_comparable'
    value = (pct * 100).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    return f'{('+' if value >= 0 else '')}{value}%'

def check_c09_movement_review(inputs, earlier_rows):
    period = inputs['period']
    expected_text = 'Warns only when a line moves by more than EUR 100.00 and by more than 20%'
    tolerance = 'EUR 100.00 and 20%'
    if any((row['status'] != 'pass' for row in earlier_rows)):
        return (summary_row('C09', period, expected_text, 'not_run - critical controls did not all pass; movement not reviewed on unreliable figures', tolerance, 'not_run', 0, False), [])
    entities = [entity for entity, _ in expected_exports(inputs)]
    prior, reason = prior_actuals_for_review(inputs, entities)
    if prior is None:
        return (summary_row('C09', period, expected_text, f'not_run - {reason}', tolerance, 'not_run', 0, False), [])
    current = build_actuals(inputs['transactions'], inputs['account_mapping'], period, entities)
    results = review_movements(current, prior, entities)
    exceptions, observed_parts = ([], [])
    for r in results:
        label = f'{r['entity']}/{r['report_line']}'
        if r['outcome'] == 'not_run':
            observed_parts.append(f'{label} not_run (no single prior or current value)')
            continue
        if r['outcome'] == 'silent':
            continue
        filename = f'transactions_{period}_{r['entity']}.csv'
        if r['outcome'] == 'new_activity':
            issue = 'new activity — no prior-period comparison'
            observed_value = f'{money_text(r['movement'], signed=True)}; movement_pct not_comparable'
        else:
            issue = f'{r['report_line']} moved beyond the review thresholds'
            observed_value = f'{money_text(r['prior'])} -> {money_text(r['current'])}; movement {money_text(r['movement'], signed=True)}; movement_pct {pct_text(r['movement_pct'])}'
        observed_parts.append(f'{label} WARNING {observed_value}')
        exceptions.append(exception_row('C09', issue, observed_value, expected_text, 'Review with the unit and explain the movement in the handover note; do not alter data to silence a warning.', source_file=filename, entity=r['entity'], period=period))
    warnings = len(exceptions)
    silent = sum((1 for r in results if r['outcome'] == 'silent'))
    not_run = sum((1 for r in results if r['outcome'] == 'not_run'))
    status = 'not_run' if not_run else 'warning' if warnings else 'pass'
    observed = f'{warnings} warning(s), {silent} silent, {not_run} not_run (prior {inputs['prior_period']})'
    if observed_parts:
        observed += ': ' + '; '.join(observed_parts)
    return (summary_row('C09', period, expected_text, observed, tolerance, status, warnings, bool(exceptions)), exceptions)

def decide_release(control_summary_rows):
    if any((r['severity'] == 'critical' and r['status'] == 'fail' for r in control_summary_rows)):
        return 'blocked'
    if any((CHECKS[r['check_id']][2] and r['status'] == 'not_run' for r in control_summary_rows)):
        return 'blocked'
    if any((r['status'] == 'warning' for r in control_summary_rows)):
        return 'ready_with_warnings'
    return 'ready'

def run_controls(inputs, learner_checks=None):
    summary_rows, exceptions = ([], [])
    for check in (check_c01_completeness_files, check_c02_completeness_rows, check_c03_duplicate_business_key, check_c04_mapping_cardinality, check_c05_unmapped_accounts, check_c06_required_values, check_c07_currency_supported, check_c08_reconciliation):
        if learner_checks and check == check_c03_duplicate_business_key:
            row, found = check(inputs, learner_checks.duplicate_rows)
        elif learner_checks and check == check_c08_reconciliation:
            row, found = check(inputs, learner_checks.reconciliation_status)
        else:
            row, found = check(inputs)
        summary_rows.append(row)
        exceptions.extend(found)
    row, found = check_c09_movement_review(inputs, summary_rows)
    summary_rows.append(row)
    exceptions.extend(found)
    control_summary = pd.DataFrame(summary_rows, columns=SUMMARY_COLUMNS)
    exception_table = pd.DataFrame(exceptions, columns=EXCEPTION_COLUMNS)
    return (control_summary, exception_table, decide_release(summary_rows))

def build_actuals(transactions, account_mapping, period, entities):
    posted = transactions[transactions['status'] == 'posted'].copy()
    if (posted['currency'] != 'EUR').any():
        raise ValueError('Posted rows in a currency other than EUR -- run controls first (C07)')
    posted['amount_value'] = parse_amount(posted['amount'])
    if posted['amount_value'].isna().any():
        raise ValueError('Posted rows with a blank or non-numeric amount -- run controls first (C06)')
    mapping = account_mapping[['account_code', 'report_line']]
    joined = posted.merge(mapping, on='account_code', how='left', validate='many_to_one', indicator=True)
    if (joined['_merge'] != 'both').any():
        raise ValueError('Posted rows with an unmapped account code -- run controls first (C05)')
    if len(joined) != len(posted):
        raise ValueError('The mapping join changed the number of rows')
    unknown_lines = set(joined['report_line']) - set(REPORT_LINES)
    if unknown_lines:
        raise ValueError(f'Mapping uses report lines outside {REPORT_LINES}: {sorted(unknown_lines)}')
    totals = joined.groupby(['entity', 'report_line'])['amount_value'].sum()
    rows = []
    for entity in entities:
        for line in REPORT_LINES:
            rows.append({'period': period, 'entity': entity, 'report_line': line, 'actual_amount': float(totals.get((entity, line), 0.0))})
    return pd.DataFrame(rows, columns=['period', 'entity', 'report_line', 'actual_amount'])

def build_summary(actuals, budget, entities_table, period):
    budget = budget.copy()
    budget['budget_amount'] = parse_amount(budget['budget_amount'])
    this_budget = budget[budget['period'] == period][['period', 'entity', 'report_line', 'budget_amount']]
    by_line = actuals.merge(this_budget, on=['period', 'entity', 'report_line'], how='left', validate='one_to_one')
    if by_line['budget_amount'].isna().any():
        missing = by_line[by_line['budget_amount'].isna()][['entity', 'report_line']].values.tolist()
        raise ValueError(f'No budget row for {period}: {missing}')
    by_line['favourable_variance'] = by_line['actual_amount'] - by_line['budget_amount']
    names = dict(zip(entities_table['entity'], entities_table['entity_name'])) if entities_table is not None else {}
    by_line.insert(2, 'entity_name', by_line['entity'].map(names).fillna(''))
    by_line.insert(4, 'report_line_label', by_line['report_line'].map(REPORT_LINE_LABELS))
    unit_rows = []
    for entity, lines in by_line.groupby('entity', sort=False):
        actual = dict(zip(lines['report_line'], lines['actual_amount']))
        planned = dict(zip(lines['report_line'], lines['budget_amount']))
        operating_result = actual['revenue'] + actual['direct_cost'] + actual['overhead']
        budget_operating_result = planned['revenue'] + planned['direct_cost'] + planned['overhead']
        unit_rows.append({'period': period, 'entity': entity, 'entity_name': names.get(entity, ''), 'revenue': actual['revenue'], 'direct_cost': actual['direct_cost'], 'overhead': actual['overhead'], 'direct_cost_expense': -1 * actual['direct_cost'], 'overhead_expense': -1 * actual['overhead'], 'operating_result': operating_result, 'budget_operating_result': budget_operating_result, 'favourable_variance': lines['favourable_variance'].sum()})
    units = pd.DataFrame(unit_rows)
    money_columns = ['revenue', 'direct_cost', 'overhead', 'direct_cost_expense', 'overhead_expense', 'operating_result', 'budget_operating_result', 'favourable_variance']
    total = {column: units[column].sum() for column in money_columns}
    check = to_cents(total['operating_result']) - to_cents(total['budget_operating_result'])
    if check != to_cents(total['favourable_variance']):
        raise ValueError('Favourable variance does not equal actual minus budget operating result')
    return {'by_line': by_line, 'units': units, 'total': total}
HEADER_FILL_RGB = '1F3864'
STATUS_FILLS = {'pass': 'C6EFCE', 'warning': 'FFEB9C', 'fail': 'FFC7CE', 'not_run': 'D9D9D9'}
MONEY_FORMAT = '#,##0.00'
COUNT_FORMAT = '0'
REFERENCE_PACK_VERSION = 'v0.1-draft'

def workbook_filename(period, run_number, release_status):
    suffix = '_BLOCKED' if release_status == 'blocked' else ''
    return f'Northbridge_Management_Pack_{period}_r{run_number:02d}{suffix}.xlsx'

def clear_rows(ws, first_row, last_column):
    for row in ws.iter_rows(min_row=first_row, max_row=max(ws.max_row, first_row), max_col=last_column):
        for cell in row:
            cell.value = None
            cell.style = 'Normal'

def read_definitions_block(ws):
    start = None
    for row_number in range(2, ws.max_row + 1):
        value = ws.cell(row_number, 1).value
        if isinstance(value, str) and value.startswith('Definitions'):
            start = row_number
            break
    if start is None:
        raise ValueError('ByLine definitions block not found -- is this the Northbridge template?')
    block = [(ws.cell(r, 1).value, copy(ws.cell(r, 1).font)) for r in range(start, ws.max_row + 1)]
    while block and block[-1][0] is None:
        block.pop()
    return block

def fill_byline_sheet(ws, by_line):
    definitions = read_definitions_block(ws)
    clear_rows(ws, first_row=2, last_column=8)
    next_row = 2
    if by_line is None:
        ws.cell(2, 1, 'No figures: the release is blocked. See the Controls and Exceptions sheets.')
        next_row = 3
    else:
        for _, line in by_line.iterrows():
            values = [line['period'], line['entity'], line['entity_name'], line['report_line'], line['report_line_label'], money(line['actual_amount']), money(line['budget_amount']), money(line['favourable_variance'])]
            for column, value in enumerate(values, start=1):
                cell = ws.cell(next_row, column, value)
                if column >= 6:
                    cell.number_format = MONEY_FORMAT
            next_row += 1
    next_row += 1
    for text, font in definitions:
        ws.cell(next_row, 1, text).font = font
        next_row += 1

def fill_controls_sheet(ws, control_summary):
    clear_rows(ws, first_row=2, last_column=len(SUMMARY_COLUMNS))
    for offset, (_, check) in enumerate(control_summary.iterrows()):
        row_number = 2 + offset
        for column, name in enumerate(SUMMARY_COLUMNS, start=1):
            ws.cell(row_number, column, check[name] if check[name] != '' else None)
        ws.cell(row_number, 9).number_format = COUNT_FORMAT
        status_cell = ws.cell(row_number, 8)
        status_cell.value = check['status']
        status_cell.fill = PatternFill('solid', fgColor=STATUS_FILLS[check['status']])

def fill_exceptions_sheet(ws, exceptions):
    clear_rows(ws, first_row=2, last_column=len(EXCEPTION_COLUMNS))
    if exceptions.empty:
        ws.cell(2, 1, 'No exceptions raised.')
        return
    for offset, (_, record) in enumerate(exceptions.iterrows()):
        for column, name in enumerate(EXCEPTION_COLUMNS, start=1):
            value = record[name]
            ws.cell(2 + offset, column, None if value == '' else value)

def fill_summary_sheet(ws, period, release_status, generated_at, pack_version, summary):
    ws['C3'], ws['C4'], ws['C5'], ws['C6'] = (period, release_status, generated_at, pack_version)
    for row_number in range(9, 13):
        for column in range(2 if row_number < 12 else 4, 10):
            ws.cell(row_number, column).value = None
    if summary is None:
        ws['B9'] = 'Figures withheld: release blocked. See the Controls, Exceptions and ReadMe sheets.'
        return
    columns = ['revenue', 'direct_cost_expense', 'overhead_expense', 'operating_result', 'budget_operating_result', 'favourable_variance']
    for offset, (_, unit) in enumerate(summary['units'].iterrows()):
        row_number = 9 + offset
        ws.cell(row_number, 2, unit['entity'])
        ws.cell(row_number, 3, unit['entity_name'])
        for column, name in enumerate(columns, start=4):
            ws.cell(row_number, column, money(unit[name])).number_format = MONEY_FORMAT
    for column, name in enumerate(columns, start=4):
        ws.cell(12, column, money(summary['total'][name])).number_format = MONEY_FORMAT

def readme_lines(inputs, release_status, control_summary, generated_at, pack_version, output_name):
    not_passed = control_summary[control_summary['status'] != 'pass']
    lines = [f'Reporting period: {inputs['period']} (prior period for movement review: {inputs['prior_period']})', f'Generated (UTC): {generated_at}', f'Case-pack version: {pack_version}', f'Output file: {output_name}', '', 'Inputs consumed (file names only):', f'  input manifest: {inputs['manifest_name']}', '  transaction exports: ' + (', '.join(inputs['received_files']) or 'none')]
    if inputs['missing_files']:
        lines.append('  listed but NOT received: ' + ', '.join(inputs['missing_files']))
    lines += [f'  account mapping: {inputs['mapping_name']}', '  reference tables: entities.csv, budget.csv, source_controls.csv, expected_files.csv', f'  prior period: {inputs['prior_source'] or 'none supplied'}', '', f'Release decision: {release_status}']
    if not_passed.empty:
        lines.append('Why: all nine controls passed.')
    else:
        lines.append('Why: these controls did not pass:')
        for _, check in not_passed.iterrows():
            lines.append(f'  {check['check_id']} {check['check_name']} = {check['status']} (affected {check['affected_count']}): {check['observed']}')
    if release_status == 'blocked':
        lines += ['What must happen next: resolve every failed or not-run control above with the unit or', 'finance (see the Exceptions sheet for the source file and row), then rerun. Do not edit', 'raw exports and do not circulate figures from this run. No figures are shown in this pack.']
    elif release_status == 'ready_with_warnings':
        lines += ['What must happen next: explain each warning on the Exceptions sheet in the handover note.', 'A warning is a request to look, not an instruction to change data. It stays visible.']
    lines += ['', 'Assumptions:', '  Signs: revenue positive, costs negative; operating result = revenue + direct cost + overhead.', '  Summary shows Direct cost and Overhead as positive expense magnitudes; ByLine is signed.', '  Favourable variance = actual - budget on signed values; positive is favourable.', '  Cancelled rows are excluded from totals and kept in the source files. No currency conversion.', '  Reconciliation tolerance EUR 0.01; movement review warns when abs(movement) > EUR 100.00 and', '  abs(movement %) > 20%. These are fictional teaching thresholds, not financial policy.', '  Summary!C14 is a formula; its value is computed by Excel when the file is opened.', '', 'To rerun: run the same script again with the same input folder and period.', '  Record the script, settings and output folder you used in your handover note.', '', 'All figures are synthetic. This pack is not an audit opinion.']
    return lines

def fill_readme_sheet(ws, lines):
    clear_rows(ws, first_row=2, last_column=1)
    for offset, text in enumerate(lines):
        if text:
            ws.cell(3 + offset, 1, text)

def write_workbook(output_folder, inputs, release_status, control_summary, exceptions, summary, run_number=1, overwrite=False, template_path=TEMPLATE_PATH, generated_at=None, pack_version=None):
    output_folder, template_path = (Path(output_folder), Path(template_path))
    generated_at = generated_at or datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
    pack_version = pack_version or inputs['pack_version'] or REFERENCE_PACK_VERSION
    if release_status == 'blocked':
        summary = None
    output_path = output_folder / workbook_filename(inputs['period'], run_number, release_status)
    if output_path.resolve() == template_path.resolve():
        raise ValueError('Refusing to write over the master template')
    if output_path.exists() and (not overwrite):
        raise FileExistsError(f'{output_path} already exists; rerun with overwrite=True to replace it')
    wb = load_workbook(template_path)
    fill_summary_sheet(wb['Summary'], inputs['period'], release_status, generated_at, pack_version, summary)
    fill_byline_sheet(wb['ByLine'], None if summary is None else summary['by_line'])
    fill_controls_sheet(wb['Controls'], control_summary)
    fill_exceptions_sheet(wb['Exceptions'], exceptions)
    fill_readme_sheet(wb['ReadMe'], readme_lines(inputs, release_status, control_summary, generated_at, pack_version, output_path.name))
    ws = wb['Summary']
    ws['B14'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[14].height = 48
    ws.merge_cells('B15:I15')
    ws['B15'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[15].height = 30
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_area = 'B2:I15'
    output_folder.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path

def write_csv(table, path):
    table.to_csv(path, index=False, lineterminator='\n', encoding='utf-8')

def main(input_folder, period, output_folder=DEFAULT_OUTPUT_FOLDER, overwrite=False, manifest_path=None, mapping_file=None, run_number=1, template_path=TEMPLATE_PATH, generated_at=None, pack_version=None):
    output_folder = Path(output_folder)
    targets = [output_folder / name for name in ('control_summary.csv', 'exceptions.csv', 'actuals_by_line.csv')]
    targets += [output_folder / workbook_filename(period, run_number, status) for status in ('ready', 'blocked')]
    existing = [path for path in targets if path.exists()]
    if existing and (not overwrite):
        raise FileExistsError(f'{existing[0]} already exists; rerun with overwrite=True to replace it')
    print(f'Northbridge reporting run -- period {period}')
    inputs = load_inputs(input_folder, period, manifest_path=manifest_path, mapping_file=mapping_file)
    print(f'  manifest: {inputs['manifest_name']}  (input folder: {Path(input_folder).as_posix()})')
    for name in inputs['listed_files']:
        rows = int((inputs['transactions']['source_file'] == name).sum())
        print(f'  loaded {name}: {rows} rows' if name in inputs['received_files'] else f'  MISSING {name}')
    print(f'  account mapping: {inputs['mapping_name']}')
    control_summary, exceptions, release_status = run_controls(inputs)
    for _, check in control_summary.iterrows():
        print(f'  {check['check_id']} {check['check_name']:<29} {check['status']:<8} affected={check['affected_count']}')
    print(f'  release: {release_status}  (exception rows: {len(exceptions)})')
    output_folder.mkdir(parents=True, exist_ok=True)
    for path in existing:
        path.unlink()
        print(f'  replacing the earlier exercise file {path.name}')
    write_csv(control_summary, output_folder / 'control_summary.csv')
    write_csv(exceptions, output_folder / 'exceptions.csv')
    written = ['control_summary.csv', 'exceptions.csv']
    summary = None
    if release_status != 'blocked':
        entities = [entity for entity, _ in expected_exports(inputs)]
        actuals = build_actuals(inputs['transactions'], inputs['account_mapping'], period, entities)
        summary = build_summary(actuals, inputs['budget'], inputs['entities'], period)
        csv_actuals = actuals.assign(actual_amount=actuals['actual_amount'].map(money_text))
        write_csv(csv_actuals, output_folder / 'actuals_by_line.csv')
        written.append('actuals_by_line.csv')
    workbook_path = write_workbook(output_folder, inputs, release_status, control_summary, exceptions, summary, run_number=run_number, overwrite=overwrite, template_path=template_path, generated_at=generated_at, pack_version=pack_version)
    written.append(workbook_path.name)
    print(f'  wrote to {output_folder.as_posix()}: {', '.join(written)}')
    return {'inputs': inputs, 'control_summary': control_summary, 'exceptions': exceptions, 'release_status': release_status, 'summary': summary, 'workbook_path': workbook_path}
